import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from bulk_mailer.core import (
    AddressBook, Contact, DEFAULT_TEMPLATE, InputError, Settings,
    load_settings, normalize_name, parse_cc, parse_contracts, prepare_mails,
    read_address_book, read_text_file, render_body, save_settings,
)


def record(number="K-001", name="山田 太郎", dates="2026/11/01 - 2027/03/31"):
    return "\n\n".join([number, dates, "サンプル受入株式会社", "東京事業所", "開発部", "システム開発", "第一チーム",
                          name, "サンプル派遣株式会社", "承認済み\n不要行1", "不要行2"])


class ParserTests(unittest.TestCase):
    def test_multiple_records_with_leading_blanks_and_crlf(self):
        source = "\ufeff\r\n　\r\n" + (record() + "\n\n" + record("K-002", "佐藤 花子") + "\n\n").replace("\n", "\r\n")
        contracts = parse_contracts(source)
        self.assertEqual(len(contracts), 2)
        self.assertEqual(contracts[0].values["管理番号"], "K-001")
        self.assertEqual(contracts[1].values["氏名"], "佐藤 花子")
        self.assertEqual(contracts[0].values["受入開始日"], "2026/11/01")
        self.assertEqual(contracts[0].values["受入終了日"], "2027/03/31")
        self.assertNotIn("不要行1", contracts[0].values.values())

    def test_missing_blank_does_not_shift_fields(self):
        with self.assertRaisesRegex(InputError, "区切りの空白行"):
            parse_contracts(record().replace("東京事業所\n\n開発部", "東京事業所\n開発部"))

    def test_missing_ignored_line_does_not_consume_next_person(self):
        source = record().replace("承認済み\n不要行1", "承認済み") + "\n\n" + record("K-002")
        with self.assertRaisesRegex(InputError, "不要行"):
            parse_contracts(source)

    def test_missing_last_required_field_is_rejected(self):
        with self.assertRaises(InputError):
            parse_contracts(record().split("承認済み")[0])

    def test_date_validation(self):
        for dates in ("2026/02/30 - 2027/03/31", "2027/11/01 - 2026/03/31", "2026-11-01 - 2027-03-31"):
            with self.subTest(dates=dates), self.assertRaises(InputError):
                parse_contracts(record(dates=dates))

    def test_empty_input(self):
        with self.assertRaisesRegex(InputError, "空"):
            parse_contracts("\n　\n")

    def test_text_encoding(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp932"):
                with self.subTest(encoding=encoding):
                    path.write_bytes(record().encode(encoding))
                    self.assertEqual(read_text_file(path), record())


class AddressTests(unittest.TestCase):
    def test_name_normalization(self):
        self.assertEqual(normalize_name("　山田　太郎 "), normalize_name("山田 太郎"))
        self.assertEqual(AddressBook([Contact("山田　太郎", "yamada@example.com")]).resolve("山田 太郎"), "yamada@example.com")

    def test_missing_and_ambiguous_names(self):
        book = AddressBook([Contact("山田 太郎", "a@example.com"), Contact("山田　太郎", "b@example.com")])
        with self.assertRaisesRegex(InputError, "重複"):
            book.resolve("山田太郎")
        with self.assertRaisesRegex(InputError, "未登録"):
            book.resolve("佐藤 花子")

    def test_cc_multiple_separators_and_deduplication(self):
        self.assertEqual(parse_cc("a@example.com; B@example.com\na@example.com、c@example.com"),
                         ("a@example.com", "B@example.com", "c@example.com"))
        self.assertEqual(parse_cc(""), ())
        with self.assertRaises(InputError):
            parse_cc("invalid address")

    def test_excel_fixed_columns_and_first_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contacts.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["氏名", "メールアドレス", "備考"])
            sheet.append(["山田 太郎", "yamada@example.com", "読み込み対象外"])
            sheet.append([None, None, None])
            book.create_sheet("ignored").append(["無視する"])
            book.save(path)
            book.close()
            self.assertEqual(read_address_book(path).resolve("山田太郎"), "yamada@example.com")

    def test_excel_reversed_or_shifted_columns_are_rejected(self):
        for header, row in ((["メールアドレス", "氏名"], ["yamada@example.com", "山田 太郎"]),
                            (["備考", "氏名", "メールアドレス"], ["", "山田 太郎", "yamada@example.com"])):
            with self.subTest(header=header), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "contacts.xlsx"
                book = Workbook()
                book.active.append(header)
                book.active.append(row)
                book.save(path)
                book.close()
                with self.assertRaisesRegex(InputError, "A1.*B1"):
                    read_address_book(path)

    def test_excel_invalid_rows_are_not_silently_dropped(self):
        for row in (["山田 太郎", "invalid"], [None, "a@example.com"], ["=A1", "a@example.com"]):
            with self.subTest(row=row), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "contacts.xlsx"
                book = Workbook()
                book.active.append(["氏名", "メールアドレス"])
                book.active.append(row)
                book.save(path)
                book.close()
                with self.assertRaisesRegex(InputError, "2行目"):
                    read_address_book(path)

    def test_excel_missing_header(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contacts.xlsx"
            book = Workbook()
            book.active.append(["名前", "アドレス"])
            book.save(path)
            book.close()
            with self.assertRaisesRegex(InputError, "1行目"):
                read_address_book(path)


class MailTests(unittest.TestCase):
    def setUp(self):
        self.contracts = parse_contracts(record())
        self.settings = Settings(subject="契約締結完了のお知らせ", cc="cc@example.com")
        self.book = AddressBook([Contact("山田 太郎", "yamada@example.com")])

    def test_exact_body_and_common_subject(self):
        item = prepare_mails(self.contracts, self.settings, self.book)[0]
        self.assertEqual(item.to, "yamada@example.com")
        self.assertEqual(item.subject, "契約締結完了のお知らせ")
        self.assertEqual(item.body, "下記契約の締結作業が完了しましたので、御連絡致します。\n\n"
                         "管理番号  :K-001\n労働者情報:山田 太郎(サンプル派遣株式会社)\n"
                         "受入先情報:東京事業所 開発部 システム開発\n"
                         "受入開始日:2026/11/01  受入終了日:2027/03/31")

    def test_unknown_or_unsafe_template_fields(self):
        for template in ("{未知}", "{氏名.__class__}", "{氏名!r}", "{氏名:>20}", "{氏名"):
            with self.subTest(template=template), self.assertRaises(InputError):
                render_body(template, self.contracts[0].values)
        self.assertEqual(render_body("{{メモ}} {氏名}", self.contracts[0].values), "{メモ} 山田 太郎")

    def test_subject_can_be_saved_empty_but_cannot_execute(self):
        settings = Settings()
        settings.validate()
        items = prepare_mails(self.contracts, settings, self.book)
        self.assertIn("山田 太郎", items[0].body)
        with self.assertRaisesRegex(InputError, "件名"):
            items[0].validate()
        with self.assertRaisesRegex(InputError, "件名"):
            settings.validate(execution=True)

    def test_missing_recipient_kept_visible_and_blocked(self):
        item = prepare_mails(self.contracts, self.settings, AddressBook())[0]
        self.assertTrue(item.included)
        self.assertIn("未登録", item.issue)
        with self.assertRaises(InputError):
            item.validate()

    def test_duplicate_management_numbers_are_blocked(self):
        contracts = parse_contracts(record() + "\n\n" + record())
        items = prepare_mails(contracts, self.settings, self.book)
        self.assertTrue(all("管理番号が重複" in item.issue for item in items))

    def test_valid_mail_needs_no_review_flag_and_repeat_execution_is_blocked(self):
        item = prepare_mails(self.contracts, self.settings, self.book)[0]
        item.validate()
        item.status = "下書き保存済み"
        with self.assertRaisesRegex(InputError, "再実行"):
            item.validate()

    def test_settings_round_trip_and_corruption(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "new" / "settings.json"
            self.assertEqual(load_settings(path), Settings())
            save_settings(self.settings, path)
            self.assertEqual(load_settings(path), self.settings)
            for content in ("{invalid", "[]", '{"subject":12}', '{"mode":"unknown"}'):
                path.write_text(content, encoding="utf-8")
                with self.subTest(content=content), self.assertRaises(InputError):
                    load_settings(path)


if __name__ == "__main__":
    unittest.main()

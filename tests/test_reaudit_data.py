"""Independent data re-audit: corruption recovery and imported literal data."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from bulk_mailer.core import (
    AddressBook, Contact, InputError, Settings, load_contacts, load_settings,
    parse_contracts, prepare_mails, read_address_book, save_contacts,
)
from test_core import record


class ReauditDataTests(unittest.TestCase):
    def test_deeply_nested_corrupt_settings_reports_input_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "settings.json"
            path.write_text("[" * 10_000 + "0" + "]" * 10_000, encoding="utf-8")
            with self.assertRaises(InputError):
                load_settings(path)

    def test_deeply_nested_corrupt_contacts_reports_input_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contacts.json"
            path.write_text("[" * 10_000 + "0" + "]" * 10_000, encoding="utf-8")
            with self.assertRaises(InputError):
                load_contacts(path)

    def test_imported_braces_remain_literal_in_rendered_mail(self):
        name = "山田 {氏名} 太郎"
        source = record(name=name).replace("サンプル受入株式会社", "会社{管理番号}")
        item = prepare_mails(
            parse_contracts(source),
            Settings(subject="共通件名", template="会社：{会社名}\n氏名：{氏名}\n{{固定文字}}"),
            AddressBook([Contact("山田{氏名}太郎", "yamada@example.com")]),
        )[0]
        item.validate()
        self.assertEqual(item.body, "会社：会社{管理番号}\n氏名：山田 {氏名} 太郎\n{固定文字}")

    def test_excel_far_ignored_column_cannot_hide_normalized_duplicate(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wide.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["氏名", "メールアドレス"])
            sheet.append(["山田 太郎", "first@example.com"])
            sheet.append(["山田\u3000太郎", "other@example.com"])
            # A formula in the last valid Excel column is outside the A/B import.
            sheet.cell(2, 16_384, "=1/0")
            workbook.save(path)
            workbook.close()
            book = read_address_book(path)
            self.assertEqual(len(book.contacts), 2)
            cache = Path(directory) / "contacts.json"
            save_contacts(book, cache, str(path))
            restored, _ = load_contacts(cache)
            with self.assertRaisesRegex(InputError, "重複"):
                restored.resolve("\t山田\u2002太郎\u00a0")

    def test_cache_replacement_failure_preserves_original_and_removes_temp(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contacts.json"
            original = AddressBook([Contact("山田太郎", "old@example.com")])
            save_contacts(original, path)
            previous = path.read_bytes()
            replacement = AddressBook([Contact("佐藤花子", "new@example.com")])
            with patch.object(Path, "replace", side_effect=PermissionError("file is locked")):
                with self.assertRaises(PermissionError):
                    save_contacts(replacement, path)
            self.assertEqual(path.read_bytes(), previous)
            self.assertEqual(load_contacts(path)[0], original)
            self.assertEqual(list(path.parent.glob("*.tmp")), [])


@unittest.skipUnless(os.environ.get("DISPLAY") or os.name == "nt", "GUI requires a display")
class ReauditDataGuiTests(unittest.TestCase):
    def assert_corrupt_cache_can_be_recovered(self, filename):
        from bulk_mailer.gui import Application

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / filename).write_text("[" * 10_000 + "0" + "]" * 10_000,
                                         encoding="utf-8")
            workbook = Workbook()
            workbook.active.append(["氏名", "メールアドレス"])
            workbook.active.append(["山田太郎", "recovered@example.com"])
            excel = root / "recovery.xlsx"
            workbook.save(excel)
            workbook.close()
            errors = []
            with patch("bulk_mailer.gui.messagebox.showwarning") as warning:
                app = Application(root, preview_mode=True)
                try:
                    app.report_callback_exception = lambda *error: errors.append(error)
                    # Allow the startup warning's actual scheduled callback to run.
                    app.after(130, app.quit)
                    app.mainloop()
                    warning.assert_called_once()
                    self.assertIn(filename, warning.call_args.args[1])
                    self.assertTrue(app.import_contacts(excel))
                    app.subject_var.set("復旧後の件名")
                    app.save_config()
                    app.source.insert("1.0", record())
                    app.create_list()
                    app.update()
                    self.assertEqual(len(app.items), 1)
                    self.assertEqual(app.items[0].to, "recovered@example.com")
                    app.items[0].validate()
                    self.assertEqual(errors, [])
                finally:
                    app.destroy()

    def test_corrupt_settings_warns_and_can_rebuild_a_valid_mail_list(self):
        self.assert_corrupt_cache_can_be_recovered("settings.json")

    def test_corrupt_contacts_warns_and_can_reimport_a_valid_address_book(self):
        self.assert_corrupt_cache_can_be_recovered("contacts.json")

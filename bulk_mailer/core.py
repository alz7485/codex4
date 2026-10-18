from __future__ import annotations

import json
import csv
import io
import os
import re
import tempfile
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from string import Formatter
from zipfile import ZipFile


DEFAULT_TEMPLATE = """下記契約の締結作業が完了しましたので、御連絡致します。

管理番号  :{管理番号}
労働者情報:{氏名}({派遣元})
受入先情報:{事業所名} {部署名} {業務名}
受入開始日:{受入開始日}  受入終了日:{受入終了日}"""
FIELDS = (
    "管理番号", "受入開始日", "受入終了日", "会社名", "事業所名", "部署名",
    "業務名", "チーム名", "氏名", "派遣元", "契約承認状態",
)
DATE_RANGE = re.compile(r"(\d{4}/\d{2}/\d{2})\s*-\s*(\d{4}/\d{2}/\d{2})")
EMAIL = re.compile(r"[^@\s;,<>]+@[^@\s;,<>]+\.[^@\s;,<>]+")
MAX_TEXT_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARACTERS = 5_000_000
MAX_EXCEL_XML_BYTES = 100 * 1024 * 1024
MAX_ADDRESS_ROWS = 100_000
MAX_CONTACTS_JSON_BYTES = 100 * 1024 * 1024


class InputError(ValueError):
    pass


def _has_unsafe_controls(text: str, allowed: str = "") -> bool:
    return any(character not in allowed and unicodedata.category(character).startswith("C")
               for character in text)


@dataclass(frozen=True)
class Contract:
    values: dict[str, str]
    line: int


def parse_contracts(text: str) -> list[Contract]:
    """Read the documented record structure without deleting internal blank lines."""
    if len(text) > MAX_TEXT_CHARACTERS:
        raise InputError("入力テキストが大きすぎます。500万文字以下に分割してください。")
    lines = text.removeprefix("\ufeff").splitlines()
    # Only EOF truncation of the final record's unused footer is optional.
    # Removing trailing blank lines makes EOF checks explicit without scanning
    # the remainder of the input for every record.
    while lines and not lines[-1].strip():
        lines.pop()
    position = 0

    def fail(message: str) -> None:
        raise InputError(f"{position + 1}行目付近: {message}")

    def value(label: str) -> str:
        nonlocal position
        if position >= len(lines) or not lines[position].strip():
            fail(f"{label}がありません。行の並びを確認してください。")
        result = lines[position].strip()
        if _has_unsafe_controls(result, "\t"):
            fail(f"{label}に制御文字または不可視文字が含まれています。")
        position += 1
        return result

    def blanks(required: bool = True, *, single: bool = False) -> None:
        nonlocal position
        start = position
        while position < len(lines) and not lines[position].strip():
            position += 1
        if required and start == position:
            fail("区切りの空白行がありません。")
        if single and position - start != 1:
            fail("必須項目間の空白行は1行にしてください。空欄の項目がないか確認してください。")

    blanks(required=False)
    contracts = []
    while position < len(lines):
        start = position + 1
        values = {"管理番号": value("管理番号")}
        blanks(single=True)
        dates = value("受入開始日・受入終了日")
        match = DATE_RANGE.fullmatch(dates)
        if not match:
            fail("日付は 2026/11/01 - 2027/03/31 の形式で入力してください。")
        begin, end = match.groups()
        try:
            if datetime.strptime(begin, "%Y/%m/%d") > datetime.strptime(end, "%Y/%m/%d"):
                raise InputError(f"{start}行目の契約: 受入終了日が開始日より前です。")
        except ValueError as error:
            if isinstance(error, InputError):
                raise
            fail("存在しない日付です。")
        values.update({"受入開始日": begin, "受入終了日": end})
        for label in ("会社名", "事業所名", "部署名", "業務名", "チーム名", "氏名", "派遣元", "契約承認状態"):
            blanks(single=True)
            values[label] = value(label)
        contracts.append(Contract(values, start))
        if position == len(lines):
            break
        value("契約承認状態の直後の不要行")
        if position == len(lines):
            break
        blanks()
        value("2つ目の不要行")
        if position < len(lines):
            blanks()
    if not contracts:
        raise InputError("入力テキストが空です。")
    return contracts


def read_text_file(path: str | Path) -> str:
    with Path(path).open("rb") as source:
        data = source.read(MAX_TEXT_BYTES + 1)
    if len(data) > MAX_TEXT_BYTES:
        raise InputError("txtファイルが大きすぎます。10MB以下に分割してください。")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError as error:
            raise InputError("UTF-16のtxtが破損しています。文字コードを確認して保存し直してください。") from error
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise InputError("文字コードを判定できません。UTF-8、UTF-16またはCP932のtxtを使用してください。")


def normalize_name(name: str) -> str:
    return "".join(unicodedata.normalize("NFKC", name).split())


def validate_address(address: str) -> str:
    address = address.strip()
    if _has_unsafe_controls(address) or not EMAIL.fullmatch(address):
        raise InputError(f"メールアドレスを確認してください: {address!r}")
    return address


def validate_body(body: str) -> None:
    if not isinstance(body, str):
        raise InputError("本文は文字列で入力してください。")
    if not body.strip():
        raise InputError("本文が空です。")
    if _has_unsafe_controls(body, "\r\n\t"):
        raise InputError("本文に制御文字または不可視文字が含まれています。")


def parse_cc(text: str) -> tuple[str, ...]:
    result: list[str] = []
    for part in re.split(r"[;,、；，\r\n]+", text):
        if part.strip():
            address = validate_address(part)
            if address.casefold() not in {item.casefold() for item in result}:
                result.append(address)
    return tuple(result)


@dataclass(frozen=True)
class Contact:
    name: str
    email: str
    staff_id: str = ""
    employee_id: str = ""


@dataclass
class AddressBook:
    contacts: list[Contact] = field(default_factory=list)

    def resolve(self, name: str) -> str:
        matches = [contact for contact in self.contacts if normalize_name(contact.name) == normalize_name(name)]
        if not matches:
            raise InputError("宛先未登録（社員情報に氏名がありません）")
        if len(matches) != 1:
            raise InputError("氏名の重複あり（宛先を手動で確定してください）")
        return matches[0].email


def read_address_book(path: str | Path) -> AddressBook:
    from openpyxl import load_workbook

    if Path(path).suffix.lower() == ".csv":
        return read_employee_csv(path)
    if Path(path).suffix.lower() != ".xlsx":
        raise InputError("社員情報は .csv 形式を使用してください。旧版の .xlsx も読み込めます。")
    try:
        with ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_EXCEL_XML_BYTES:
                raise InputError("Excelの展開後サイズが大きすぎます。100MB以下に分割してください。")
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except InputError:
        raise
    except Exception as error:
        raise InputError(f"Excelを読み込めません: {error}") from error
    try:
        sheet = book.worksheets[0]
        # Exported workbooks can have stale dimensions. Trusting them silently
        # drops rows, including duplicate names that must block resolution.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows, ())]
        if header[:2] != ["氏名", "メールアドレス"]:
            raise InputError("先頭シートの1行目はA1に「氏名」、B1に「メールアドレス」を入力してください。2行目以降のA列・B列を読み込みます。")
        name_column, email_column = 0, 1
        contacts = []
        for number, row in enumerate(rows, start=2):
            if number > MAX_ADDRESS_ROWS + 1:
                raise InputError("Excelの行数が多すぎます。見出しを除き10万行以下に分割してください。")
            name = str((row[name_column] if name_column < len(row) else None) or "").strip()
            address = str((row[email_column] if email_column < len(row) else None) or "").strip()
            if not name and not address:
                continue
            if not name or name.startswith("=") or address.startswith("="):
                raise InputError(f"Excelの{number}行目: 氏名・メールアドレスを数式ではなく文字列で入力してください。")
            try:
                contacts.append(Contact(name, validate_address(address)))
            except InputError as error:
                raise InputError(f"Excelの{number}行目: {error}") from error
        if not contacts:
            raise InputError("Excelに宛先が登録されていません。")
        return AddressBook(contacts)
    except InputError:
        raise
    except Exception as error:
        # read_only workbooks parse worksheet XML lazily during iteration.
        raise InputError(f"Excelを読み込めません: {error}") from error
    finally:
        book.close()


def read_employee_csv(path: str | Path) -> AddressBook:
    """Project A/B/C/J only; IDs stay strings, unused CSV columns are discarded."""
    try:
        with Path(path).open("rb") as source:
            data = source.read(MAX_CONTACTS_JSON_BYTES + 1)
        if len(data) > MAX_CONTACTS_JSON_BYTES:
            raise InputError("社員情報CSVが大きすぎます。100MiB以下に分割してください。")
        if data.startswith((b"\xff\xfe", b"\xfe\xff")):
            content = data.decode("utf-16")
        else:
            try:
                content = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                content = data.decode("cp932")
    except UnicodeDecodeError as error:
        raise InputError("社員情報CSVの文字コードを確認してください。UTF-8、UTF-16（BOMあり）、CP932に対応します。") from error
    previous_limit = csv.field_size_limit()
    try:
        csv.field_size_limit(MAX_CONTACTS_JSON_BYTES)
        rows = csv.reader(io.StringIO(content, newline=""), strict=True)
        header = next(rows, [])
        if len(header) < 10 or header[1].strip() != "派遣元スタッフコード":
            raise InputError("社員情報CSVは1行目を見出しとし、B列の見出しを「派遣元スタッフコード」にしてください。A列＝スタッフID、B列＝社員ID、C列＝氏名、J列＝メールアドレスで読み込みます。")
        contacts = []
        for number, row in enumerate(rows, start=2):
            if number > MAX_ADDRESS_ROWS + 1:
                raise InputError("社員情報CSVは見出しを除き10万行以下に分割してください。")
            selected = [row[index] if index < len(row) else "" for index in (0, 1, 2, 9)]
            if not any(value.strip() for value in selected):
                continue
            if len(row) < 10:
                raise InputError(f"CSVの{number}件目: J列までの列が不足しています。")
            staff_id, employee_id, name, email = (value.strip() for value in selected)
            if not name or _has_unsafe_controls(name, "\t"):
                raise InputError(f"CSVの{number}件目: 氏名が空欄、または不正な文字を含んでいます。")
            if any(_has_unsafe_controls(value) for value in (staff_id, employee_id)):
                raise InputError(f"CSVの{number}件目: IDに不正な文字が含まれています。")
            try:
                contacts.append(Contact(name, validate_address(email), staff_id, employee_id))
            except InputError as error:
                raise InputError(f"CSVの{number}件目: {error}") from error
        if not contacts:
            raise InputError("CSVに社員情報が登録されていません。")
        return AddressBook(contacts)
    except csv.Error as error:
        raise InputError(f"社員情報CSVの形式を確認してください: {error}") from error
    finally:
        csv.field_size_limit(previous_limit)


def render_body(template: str, values: dict[str, str]) -> str:
    try:
        for _, name, spec, conversion in Formatter().parse(template):
            if name is None:
                continue
            if name not in FIELDS or spec or conversion:
                raise InputError(f"使用できない差し込み項目です: {{{name}}}")
        return template.format_map(values)
    except (ValueError, KeyError) as error:
        if isinstance(error, InputError):
            raise
        raise InputError("本文の差し込み記号を確認してください。例: {氏名}。通常の波括弧は {{ と }} で記載できます。") from error


@dataclass
class Settings:
    subject: str = ""
    cc: str = ""
    template: str = DEFAULT_TEMPLATE
    mode: str = "draft"
    address_book_path: str = ""

    def validate(self, *, execution: bool = False) -> None:
        if not all(isinstance(value, str) for value in asdict(self).values()):
            raise InputError("設定値は文字列である必要があります。")
        if self.mode not in ("draft", "send"):
            raise InputError("実行方法を選択してください。")
        if "\r" in self.subject or "\n" in self.subject:
            raise InputError("件名には改行を使用できません。")
        if _has_unsafe_controls(self.subject):
            raise InputError("件名に制御文字または不可視文字が含まれています。")
        if execution and not self.subject.strip():
            raise InputError("設定画面で共通の件名を入力してください。")
        if not self.template.strip():
            raise InputError("本文テンプレートが空です。")
        if _has_unsafe_controls(self.template, "\r\n\t"):
            raise InputError("本文テンプレートに制御文字または不可視文字が含まれています。")
        parse_cc(self.cc)
        render_body(self.template, dict.fromkeys(FIELDS, "確認"))


def data_directory() -> Path:
    base = Path(os.environ.get("LOCALAPPDATA", Path.home() / ".local" / "share"))
    return base / "OutlookBulkMailer"


def load_settings(path: Path) -> Settings:
    if not path.exists():
        return Settings()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("設定ファイルの形式が不正です。")
        settings = Settings(**{key: data[key] for key in Settings.__dataclass_fields__ if key in data})
        if not all(isinstance(value, str) for value in asdict(settings).values()):
            raise ValueError("設定値は文字列である必要があります。")
        settings.validate()
        return settings
    except (ValueError, TypeError, OSError, RecursionError) as error:
        raise InputError(f"設定ファイルを読み込めません: {path}\n{error}") from error


def _write_json_atomic(payload: dict, path: Path, max_bytes: int | None = None) -> None:
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    if max_bytes is not None and len(text.encode("utf-8")) > max_bytes:
        raise InputError("社員情報JSONが大きすぎます。CSVを分割して読み込み直してください。")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                         prefix=path.name + ".", suffix=".tmp", delete=False) as output:
            temporary = Path(output.name)
            output.write(text)
            output.flush()
            os.fsync(output.fileno())
        temporary.replace(path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def save_settings(settings: Settings, path: Path) -> None:
    settings.validate()
    _write_json_atomic(asdict(settings), path)


def _parse_contacts_payload(payload: dict) -> tuple[AddressBook, str]:
    if not isinstance(payload, dict) or type(payload.get("version")) is not int or payload["version"] not in (1, 2):
        raise InputError("社員情報JSONの形式が不正です。CSVから読み込み直してください。")
    version = payload["version"]
    records = payload.get("employees" if version == 2 else "contacts")
    source_file = payload.get("source_file" if version == 2 else "source_excel")
    if not isinstance(records, list) or not isinstance(source_file, str):
        raise InputError("社員情報JSONの形式が不正です。CSVから読み込み直してください。")
    if _has_unsafe_controls(source_file):
        raise InputError("社員情報JSONの元ファイルパスに不正な文字が含まれています。")
    if len(records) > MAX_ADDRESS_ROWS:
        raise InputError("宛先JSONの件数が多すぎます。10万件以下にしてください。")
    contacts = []
    estimated_bytes = len(source_file.encode("utf-8")) + 128
    for number, row in enumerate(records, 1):
        if (not isinstance(row, dict) or not isinstance(row.get("name"), str)
                or not isinstance(row.get("email"), str)):
            raise InputError(f"宛先JSONの{number}件目: 氏名とメールアドレスは文字列で指定してください。")
        name = row["name"].strip()
        if not name or _has_unsafe_controls(name, "\t"):
            raise InputError(f"宛先JSONの{number}件目: 氏名が空欄、または不正な文字を含んでいます。")
        try:
            staff_id = row.get("staff_id", "") if version == 1 else row.get("staff_id")
            employee_id = row.get("employee_id", "") if version == 1 else row.get("employee_id")
            if not isinstance(staff_id, str) or not isinstance(employee_id, str):
                raise InputError("スタッフIDと社員IDは文字列で指定してください。")
            if any(_has_unsafe_controls(value) for value in (staff_id, employee_id)):
                raise InputError("IDに不正な文字が含まれています。")
            address = validate_address(row["email"])
            estimated_bytes += sum(len(value.encode("utf-8")) for value in (name, address, staff_id, employee_id)) + 128
            if estimated_bytes > MAX_CONTACTS_JSON_BYTES:
                raise InputError("社員情報JSONが大きすぎます。CSVを分割して読み込み直してください。")
            contacts.append(Contact(name, address, staff_id, employee_id))
        except InputError as error:
            raise InputError(f"宛先JSONの{number}件目: {error}") from error
    return AddressBook(contacts), source_file


def save_contacts(book: AddressBook, path: Path, source_excel: str = "") -> None:
    """Replace the complete address book only after validation and durable write."""
    payload = {"version": 2, "source_file": source_excel,
               "employees": [asdict(contact) for contact in book.contacts]}
    checked, _ = _parse_contacts_payload(payload)
    payload["employees"] = [asdict(contact) for contact in checked.contacts]
    _write_json_atomic(payload, path, MAX_CONTACTS_JSON_BYTES)


def load_contacts(path: Path) -> tuple[AddressBook, str]:
    try:
        with path.open("rb") as source:
            data = source.read(MAX_CONTACTS_JSON_BYTES + 1)
        if len(data) > MAX_CONTACTS_JSON_BYTES:
            raise InputError("社員情報JSONが大きすぎます。CSVを分割して読み込み直してください。")
        return _parse_contacts_payload(json.loads(data.decode("utf-8")))
    except (ValueError, TypeError, OSError, RecursionError) as error:
        raise InputError(f"宛先JSONを読み込めません: {path}\n{error}") from error


@dataclass
class MailItem:
    contract: Contract
    to: str
    cc: tuple[str, ...]
    subject: str
    body: str
    issue: str = ""
    included: bool = True
    status: str = "未実行"

    def validate_content(self) -> None:
        if self.issue:
            raise InputError(self.issue)
        validate_address(self.to)
        for address in self.cc:
            validate_address(address)
        if not self.subject.strip():
            raise InputError("設定画面で共通の件名を入力してください。")
        if _has_unsafe_controls(self.subject):
            raise InputError("件名に制御文字・改行または不可視文字が含まれています。")
        validate_body(self.body)

    def validate(self) -> None:
        self.validate_content()
        if self.status != "未実行":
            raise InputError("実行済み、または実行結果が未確定のメールは再実行できません。")


def prepare_mails(contracts: list[Contract], settings: Settings, book: AddressBook) -> list[MailItem]:
    settings.validate()
    # Build an index for this preparation only: the public contacts list is
    # mutable, so a persistent cache could silently resolve an outdated address.
    recipients: dict[str, list[str]] = {}
    for contact in book.contacts:
        recipients.setdefault(normalize_name(contact.name), []).append(contact.email)
    cc = parse_cc(settings.cc)
    counts: dict[str, int] = {}
    for contract in contracts:
        key = contract.values["管理番号"]
        counts[key] = counts.get(key, 0) + 1
    result = []
    for contract in contracts:
        issue = ""
        matches = recipients.get(normalize_name(contract.values["氏名"]), [])
        if not matches:
            to, issue = "", "宛先未登録（社員情報に氏名がありません）"
        elif len(matches) != 1:
            to, issue = "", "氏名の重複あり（宛先を手動で確定してください）"
        else:
            to = matches[0]
        if counts[contract.values["管理番号"]] > 1:
            issue = "管理番号が重複しています。入力テキストを確認してください。"
        result.append(MailItem(contract, to, cc, settings.subject,
                               render_body(settings.template, contract.values), issue))
    return result
